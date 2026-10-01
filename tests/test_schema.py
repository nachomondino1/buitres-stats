"""Valida que Buitres_v3.xlsx tenga las hojas y columnas que define schema.py.

Corre contra data/Buitres_v3.xlsx (no versionado, se genera con
`python scripts/build_v3.py --build`). Si no existe, se salta.
"""
from pathlib import Path

import openpyxl
import pytest

import schema

V3_PATH = Path(__file__).parent.parent / "data" / "Buitres_v3.xlsx"

pytestmark = pytest.mark.skipif(
    not V3_PATH.exists(), reason="data/Buitres_v3.xlsx no existe (correr scripts/build_v3.py --build)"
)


@pytest.fixture(scope="module")
def wb():
    return openpyxl.load_workbook(V3_PATH, data_only=False)


def encabezados(ws):
    valores = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    # convención v2/v3: columnas fórmula llevan sufijo visible " (auto)"
    return [v[: -len(" (auto)")] if isinstance(v, str) and v.endswith(" (auto)") else v for v in valores]


def test_hojas_esperadas(wb):
    for hoja in schema.HOJAS_V3:
        assert hoja in wb.sheetnames, f"falta la hoja {hoja!r}"


def test_no_quedan_hojas_de_visualizacion(wb):
    for hoja in ("stats_jugadores", "stats_equipo"):
        assert hoja not in wb.sheetnames, f"{hoja!r} debería haberse borrado en v3"


@pytest.mark.parametrize("hoja", list(schema.COLUMNAS.keys()))
def test_columnas_esperadas(wb, hoja):
    ws = wb[hoja]
    esperadas = schema.COLUMNAS[hoja] + schema.COLUMNAS_QC.get(hoja, [])
    assert encabezados(ws) == esperadas, f"columnas de {hoja!r} no coinciden con schema.py"


def test_columnas_internas_siguen_en_partidos(wb):
    ws = wb["partidos"]
    headers = encabezados(ws)
    for col in schema.COLUMNAS_INTERNAS["partidos"]:
        assert col in headers, f"'{col}' debería seguir en partidos (no se publica, pero no se borra)"


def test_en_filtro_borradas(wb):
    for hoja in ("partidos", "alineaciones", "goles"):
        headers = [c.value for c in next(wb[hoja].iter_rows(min_row=1, max_row=1))]
        assert not any(h and "en_filtro" in str(h) for h in headers), f"{hoja!r} todavía tiene en_filtro"


def test_nombre_mostrar_unico(wb):
    ws = wb["jugadores"]
    filas = list(ws.iter_rows(min_row=2, values_only=True))
    nombres, apellidos, apodos, origenes = (
        [r[1] for r in filas], [r[2] for r in filas], [r[3] for r in filas], [r[4] for r in filas],
    )
    previews = []
    for nombre, apellido, apodo, origen in zip(nombres, apellidos, apodos, origenes):
        base = f"{nombre} {apellido}".strip() if apellido else nombre
        if apodo:
            base += f' "{apodo}"'
        if origen and not apellido:
            base += f" ({origen})"
        previews.append(base)
    assert len(previews) == len(set(previews)), "nombre_mostrar duplicado"
