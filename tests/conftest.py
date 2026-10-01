"""Fixture builder para los tests de export_data.py: arma un Buitres_v3.xlsx
sintético mínimo en memoria (no el dato real, que no va al repo público). Cada
test parte de la base "limpia" (sin ERROR ni WARNING) y pisa solo lo que
necesita para disparar una regla puntual.
"""
import datetime as dt
from pathlib import Path

import openpyxl
import pytest

# 7 jugadores (mínimo de alineación válido) con nombre_mostrar == nombre (sin
# apellido/apodo/origen) para que los tests no tengan que pensar en la fórmula.
JUGADORES_BASE = [
    {"id_jugador": "J01", "nombre": "Ana", "apellido": None, "apodo": None, "origen": None},
    {"id_jugador": "J02", "nombre": "Bruno", "apellido": None, "apodo": None, "origen": None},
    {"id_jugador": "J03", "nombre": "Carla", "apellido": None, "apodo": None, "origen": None},
    {"id_jugador": "J04", "nombre": "Diego", "apellido": None, "apodo": None, "origen": None},
    {"id_jugador": "J05", "nombre": "Elena", "apellido": None, "apodo": None, "origen": None},
    {"id_jugador": "J06", "nombre": "Fabio", "apellido": None, "apodo": None, "origen": None},
    {"id_jugador": "J07", "nombre": "Gina", "apellido": None, "apodo": None, "origen": None},
]

PARTIDOS_BASE = [
    {
        "id_partido": 1, "fecha": dt.date(2026, 1, 10), "hora": None, "tipo": "Amistoso",
        "rival": "Rival A", "gf_1t": 1, "gc_1t": 0, "gf_2t": 0, "gc_2t": 0,
        "link_video": None, "lavo_camisetas": None, "observaciones": None,
    },
]

ALINEACIONES_BASE = [
    {"id_partido": 1, "jugador": nombre, "amarillas": 0, "rojas": 0}
    for nombre in ("Ana", "Bruno", "Carla", "Diego", "Elena", "Fabio", "Gina")
]

GOLES_BASE = [
    {
        "id_partido": 1, "tipo_gol": "GF", "nro_gol": 1, "goleador": "Ana", "asistidor": "Bruno",
        "fuente": "Contraataque", "marcador_tras_gol": "1-0", "link": "https://example.com/gol1",
    },
]

LISTAS_BASE = {
    "tipo_partido": ["Torneo", "Amistoso"],
    "fuente_gol": ["Sin dato", "Contraataque", "Gol en contra"],
    "tipo_gol": ["GF", "GC"],
}

COLUMNAS = {
    "jugadores": ["id_jugador", "nombre", "apellido", "apodo", "origen"],
    "partidos": [
        "id_partido", "fecha", "hora", "tipo", "rival",
        "gf_1t", "gc_1t", "gf_2t", "gc_2t", "link_video", "lavo_camisetas", "observaciones",
    ],
    "alineaciones": ["id_partido", "jugador", "amarillas", "rojas"],
    "goles": ["id_partido", "tipo_gol", "nro_gol", "goleador", "asistidor", "fuente", "marcador_tras_gol", "link"],
    "listas": ["tipo_partido", "fuente_gol", "tipo_gol"],
}


def _escribir_hoja(wb, nombre, columnas, filas):
    ws = wb.create_sheet(nombre)
    ws.append(columnas)
    for fila in filas:
        ws.append([fila.get(c) for c in columnas])


def build_workbook(path, *, jugadores=None, partidos=None, alineaciones=None, goles=None, listas=None,
                    quitar_hoja=None, quitar_columna=None):
    """Escribe en `path` un Buitres_v3.xlsx sintético. Cada tabla por default usa
    la base limpia; pasar una lista propia la reemplaza entera. `quitar_hoja`
    omite esa hoja (para el caso 'falta la hoja'); `quitar_columna` es
    (hoja, columna) para el caso 'falta la columna'."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    tablas = {
        "jugadores": jugadores if jugadores is not None else JUGADORES_BASE,
        "partidos": partidos if partidos is not None else PARTIDOS_BASE,
        "alineaciones": alineaciones if alineaciones is not None else ALINEACIONES_BASE,
        "goles": goles if goles is not None else GOLES_BASE,
    }

    for hoja, filas in tablas.items():
        if hoja == quitar_hoja:
            continue
        cols = list(COLUMNAS[hoja])
        if quitar_columna and quitar_columna[0] == hoja:
            cols.remove(quitar_columna[1])
        _escribir_hoja(wb, hoja, cols, filas)

    if "listas" != quitar_hoja:
        listas_dict = listas if listas is not None else LISTAS_BASE
        ws = wb.create_sheet("listas")
        ws.append(COLUMNAS["listas"])
        maxlen = max(len(v) for v in listas_dict.values())
        for i in range(maxlen):
            ws.append([listas_dict[c][i] if i < len(listas_dict[c]) else None for c in COLUMNAS["listas"]])

    wb.save(path)
    return path


@pytest.fixture
def workbook_path(tmp_path):
    def _build(**kwargs):
        return build_workbook(tmp_path / "fixture.xlsx", **kwargs)
    return _build
