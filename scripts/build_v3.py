"""
Genera Buitres_v3.xlsx a partir de Buitres_v2.xlsx (reproducible, sin edición manual).

Paso 1: separa jugadores.nombre en nombre/apellido/apodo/origen con heurísticas
simples y escribe scripts/propuesta_jugadores.csv para que la usuaria lo revise.

Paso 2 (--build, una vez confirmado el CSV): construye Buitres_v3.xlsx completo
(borra stats_jugadores/stats_equipo, borra columnas en_filtro, reconstruye
jugadores con nombre_mostrar como fórmula, reindexa checks y data validations).

Uso:
    python scripts/build_v3.py                  # (re)genera la propuesta, no toca nada más
    python scripts/build_v3.py --build           # construye Buitres_v3.xlsx desde el CSV confirmado
"""
import csv
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, str(Path(__file__).parent.parent))
import schema

ENTRADA = "data/Buitres_v2.xlsx"
SALIDA_PROPUESTA = Path(__file__).parent / "propuesta_jugadores.csv"
SALIDA_V3 = "data/Buitres_v3.xlsx"

COLOR_HEADER_MANUAL = "FF1F3864"  # azul oscuro: se carga a mano (convención de v2)
COLOR_HEADER_FORMULA = "FF7F7F7F"  # gris: fórmula, no tocar
COLOR_FONT_HEADER = "FFFFFFFF"

# Palabras que delatan un descriptor de posición/rol en vez de un apellido real
# (no hay forma de distinguirlo del apellido solo con texto: se marca ambiguo).
DESCRIPTORES_ROL = {"arquero", "defensor", "delantero", "mediocampista", "capitan", "capitán", "profe", "dt"}


def leer_nombres(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb["jugadores"]
    filas = list(ws.iter_rows(min_row=2, values_only=True))
    return [(id_jugador, nombre) for id_jugador, nombre in filas if nombre is not None]


def palabras_entre_parentesis(nombres):
    """Vocabulario de orígenes ya explícitos (entre paréntesis), para detectar
    apellidos que en realidad son el mismo descriptor sin paréntesis (caso
    "Roger Fabrica" vs "Tomas (fabrica)")."""
    vocab = set()
    for _, nombre in nombres:
        for m in re.finditer(r"\(([^)]+)\)", nombre):
            vocab.update(p.lower() for p in m.group(1).split())
    return vocab


def separar_nombre(nombre_crudo, vocab_origenes):
    texto = nombre_crudo
    origen = None
    apodo = None
    ambiguo = False
    motivo = None

    m = re.search(r"\(([^)]+)\)", texto)
    if m:
        origen = m.group(1).strip()
        texto = (texto[: m.start()] + texto[m.end() :]).strip()

    m2 = re.search(r'"([^"]+)"', texto)
    if m2:
        apodo = m2.group(1).strip()
        texto = re.sub(r"\s+", " ", (texto[: m2.start()] + texto[m2.end() :])).strip()
        ambiguo = True
        motivo = "apodo entre comillas: confirmar separación nombre/apellido"

    palabras = texto.split()
    if not palabras:
        nombre, apellido = nombre_crudo, None
        ambiguo = True
        motivo = "no se pudo separar nombre/apellido"
    elif len(palabras) == 1:
        nombre, apellido = palabras[0], None
        if origen is None and apodo is None:
            ambiguo = True
            motivo = "una sola palabra: ¿alcanza con nombre o falta apellido?"
    else:
        nombre = palabras[0]
        apellido = " ".join(palabras[1:])
        if origen is None and apodo is None and palabras[-1].lower() in (DESCRIPTORES_ROL | vocab_origenes):
            ambiguo = True
            motivo = f"'{palabras[-1]}' parece descriptor/origen, no apellido"

    return nombre, apellido, apodo, origen, ambiguo, motivo


def generar_propuesta(entrada=ENTRADA, salida=SALIDA_PROPUESTA):
    nombres = leer_nombres(entrada)
    vocab = palabras_entre_parentesis(nombres)

    filas = []
    for id_jugador, nombre_crudo in nombres:
        nombre, apellido, apodo, origen, ambiguo, motivo = separar_nombre(nombre_crudo, vocab)
        filas.append({
            "id_jugador": id_jugador,
            "nombre_actual": nombre_crudo,
            "nombre": nombre,
            "apellido": apellido or "",
            "apodo": apodo or "",
            "origen": origen or "",
            "nombre_mostrar_preview": schema.nombre_mostrar(nombre, apellido, apodo, origen),
            "ambiguo": "SI" if ambiguo else "",
            "motivo": motivo or "",
        })

    with open(salida, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(filas[0].keys()))
        w.writeheader()
        w.writerows(filas)

    ambiguos = [f for f in filas if f["ambiguo"]]
    print(f"{len(filas)} jugadores -> {salida}")
    print(f"{len(ambiguos)} marcados como ambiguos (revisar antes de confirmar):")
    for f in ambiguos:
        print(f"  - {f['id_jugador']} {f['nombre_actual']!r}: {f['motivo']}")
    print("\nRevisá y corregí el CSV a mano si hace falta. Recién con el CSV confirmado"
          " se construye la hoja 'jugadores' definitiva de Buitres_v3.xlsx (paso 2, pendiente).")


def _leer_propuesta_confirmada(path):
    with open(path, newline="", encoding="utf-8") as f:
        filas = list(csv.DictReader(f))
    jugadores = []
    for r in filas:
        jugadores.append({
            "id_jugador": r["id_jugador"],
            "nombre": r["nombre"],
            "apellido": r["apellido"] or None,
            "apodo": r["apodo"] or None,
            "origen": r["origen"] or None,
        })
    previews = [schema.nombre_mostrar(j["nombre"], j["apellido"], j["apodo"], j["origen"]) for j in jugadores]
    dups = {p for p in previews if previews.count(p) > 1}
    if dups:
        raise SystemExit(f"nombre_mostrar duplicado en la propuesta confirmada, no se puede construir v3: {dups}")
    return jugadores


def _estilar_header(ws, col, formula):
    cell = ws.cell(row=1, column=col)
    color = COLOR_HEADER_FORMULA if formula else COLOR_HEADER_MANUAL
    cell.fill = PatternFill(fgColor=color, fill_type="solid")
    cell.font = Font(color=COLOR_FONT_HEADER, bold=True)


def _reconstruir_jugadores(wb, jugadores):
    ws = wb["jugadores"]
    # limpiar hoja entera (incluye la vieja columna 'nombre' de v2) y reescribir
    ws.delete_rows(1, ws.max_row)
    headers = ["id_jugador", "nombre", "apellido", "apodo", "origen", "nombre_mostrar (auto)"]
    for col, h in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=h)
        _estilar_header(ws, col, formula=(h == "nombre_mostrar (auto)"))
    for i, j in enumerate(jugadores, start=2):
        ws.cell(row=i, column=1, value=j["id_jugador"])
        ws.cell(row=i, column=2, value=j["nombre"])
        ws.cell(row=i, column=3, value=j["apellido"])
        ws.cell(row=i, column=4, value=j["apodo"])
        ws.cell(row=i, column=5, value=j["origen"])
        ws.cell(row=i, column=6, value=(
            f'=TRIM(B{i}&" "&C{i})'
            f'&IF(D{i}<>""," "&""""&D{i}&"""","")'
            f'&IF(AND(E{i}<>"",C{i}=""),"' f' ("&E{i}&")","")'
        ))
    ws.freeze_panes = "A2"
    anchos = {1: 12, 2: 14, 3: 18, 4: 14, 5: 24, 6: 30}
    for col, w in anchos.items():
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = w


def _reindexar_checks_goles(wb):
    ws_goles = wb["goles"]
    ws_checks = wb["checks"]
    # check_dup y check_jugó pasan de J,K a I,J al borrar en_filtro (col I)
    ws_checks["B4"] = '=COUNTIF(goles!$I$2:$I$600,"DUP")'
    ws_checks["B5"] = '=COUNTIF(goles!$J$2:$J$600,"NO JUGÓ")'

    vieja = ws_goles.conditional_formatting["J2:K600"]
    fill_vieja = vieja[0].dxf.fill
    del ws_goles.conditional_formatting["J2:K600"]
    ws_goles.conditional_formatting.add("I2:J600", FormulaRule(formula=['I2<>""'], fill=fill_vieja))


def _repuntar_dropdowns_jugadores(wb):
    # las listas de jugador apuntaban a jugadores!$B (nombre); ahora el desplegable
    # muestra nombre_mostrar, que quedó en la columna F
    objetivos = [("partidos", "N2:N200"), ("alineaciones", "B2:B600"), ("goles", "D2:E600")]
    for hoja, rango in objetivos:
        ws = wb[hoja]
        for dv in ws.data_validations.dataValidation:
            if dv.sqref == rango and dv.formula1 and dv.formula1.startswith("jugadores!$B$"):
                dv.formula1 = dv.formula1.replace("$B$", "$F$")


def _actualizar_leeme(wb):
    ws = wb["Leeme"]
    filas_originales = ws.max_row
    textos = [c.value for c in ws["A"]]
    nuevo = []
    for t in textos:
        if t is None:
            nuevo.append("")  # separador en blanco; "" en vez de None (ver nota más abajo)
        elif t.startswith("Buitres"):
            nuevo.append("Buitres – estadísticas v3")
        elif "Celdas amarillas en stats_jugadores" in t:
            continue
        elif t.startswith("No ordenes las hojas stats_*"):
            continue
        elif t.startswith("La clave del jugador es su nombre"):
            nuevo.append(
                "La clave estable del jugador es id_jugador. El desplegable muestra "
                "nombre_mostrar (nombre + apellido + \"apodo\" + origen si no hay apellido); "
                "no lo edites a mano, es fórmula."
            )
        elif t.startswith("Datos migrados con migrar_buitres.py"):
            nuevo.append(
                "Datos migrados con scripts/archivo/migrar_buitres.py. El gol de Roger Fabrica con "
                "marcador 1-4 quedó corregido: es del 26/09/2026 (ya aplicado). Ese partido "
                "(26/09 vs Perez el ratón) todavía tiene los goles en contra (GC) sin cargar: "
                "checks lo marca incompleto hasta que se carguen."
            )
        else:
            nuevo.append(t)
    # ws.cell(value=None) es un no-op en openpyxl (None = "no cambiar"), no vacía la celda;
    # por eso se asigna cell.value directo para poder borrar las filas sobrantes.
    for r in range(1, filas_originales + 1):
        ws.cell(row=r, column=1).value = nuevo[r - 1] if r - 1 < len(nuevo) else None


def _mover_gol_roger_fabrica(wb):
    """Corrección de datos confirmada por la usuaria (documentada en DECISIONS.md):
    el gol de Roger Fabrica (asistió Gonzalo Ruiz Diaz, fuente Lateral, marcador 1-4)
    estaba cargado en id_partido=16 (19/09, Celta FC) con nro_gol duplicado; es en
    realidad del id_partido=17 (26/09, Perez el ratón), que no tenía goles cargados.
    A diferencia de lo que decía la spec, Buitres_v2.xlsx NO traía esto corregido:
    se detectó al verificar los checks (ver DECISIONS.md)."""
    ws = wb["goles"]
    movidas = 0
    for row in ws.iter_rows(min_row=2):
        if row[0].value == 16 and row[1].value == "GF" and row[3].value == "Roger Fabrica":
            row[0].value = 17
            movidas += 1
    if movidas != 1:
        raise SystemExit(
            f"Se esperaba mover exactamente 1 fila de gol (Roger Fabrica, partido 16->17); "
            f"se encontraron {movidas}. Revisar a mano, no se construyó Buitres_v3.xlsx."
        )


def _renombrar_referencias_jugadores(wb, jugadores):
    """Si separar nombre/apellido/apodo/origen cambió el texto que identifica a un
    jugador (p.ej. 'Bona' -> 'Tomas Bonamino'), hay que propagar ese cambio adonde
    ese texto viejo está tipeado a mano: alineaciones.jugador, goles.goleador/
    asistidor y partidos.lavo_camisetas (el 'Buscar y reemplazar' que pedía el
    Leeme de v2 cuando un jugador cambia de nombre)."""
    ws_j = wb["jugadores"]
    nombre_viejo_por_id = {
        ws_j.cell(row=r, column=1).value: ws_j.cell(row=r, column=2).value
        for r in range(2, ws_j.max_row + 1)
        if ws_j.cell(row=r, column=1).value is not None
    }
    mapa = {}
    for j in jugadores:
        viejo = nombre_viejo_por_id.get(j["id_jugador"])
        nuevo = schema.nombre_mostrar(j["nombre"], j["apellido"], j["apodo"], j["origen"])
        if viejo and viejo != nuevo:
            mapa[viejo] = nuevo

    if not mapa:
        return mapa

    objetivos = [("alineaciones", 2), ("goles", 4), ("goles", 5), ("partidos", 14)]
    for hoja, col in objetivos:
        ws = wb[hoja]
        for r in range(2, ws.max_row + 1):
            celda = ws.cell(row=r, column=col)
            if celda.value in mapa:
                celda.value = mapa[celda.value]
    return mapa


def construir_v3(entrada=ENTRADA, propuesta_csv=SALIDA_PROPUESTA, salida=SALIDA_V3):
    jugadores = _leer_propuesta_confirmada(propuesta_csv)

    wb = openpyxl.load_workbook(entrada, data_only=False)

    _mover_gol_roger_fabrica(wb)
    renombrados = _renombrar_referencias_jugadores(wb, jugadores)
    if renombrados:
        print("Nombres propagados (Buscar y reemplazar) en alineaciones/goles/partidos:")
        for viejo, nuevo in renombrados.items():
            print(f"  - {viejo!r} -> {nuevo!r}")

    for hoja in ("stats_jugadores", "stats_equipo"):
        if hoja in wb.sheetnames:
            del wb[hoja]

    wb["partidos"].delete_cols(16, 1)      # P: en_filtro (auto)
    wb["alineaciones"].delete_cols(5, 1)   # E: en_filtro (auto)
    wb["goles"].delete_cols(9, 1)          # I: en_filtro (auto)

    _reindexar_checks_goles(wb)
    _repuntar_dropdowns_jugadores(wb)
    _reconstruir_jugadores(wb, jugadores)
    _actualizar_leeme(wb)

    wb.save(salida)
    print(f"-> {salida}")
    print(f"{len(jugadores)} jugadores en la hoja 'jugadores' nueva.")
    print("Pendiente (paso manual, fuera de este script): abrir en Excel real para que")
    print("recalcule, y confirmar 0 errores de fórmula antes de considerar la Fase 1 cerrada.")


if __name__ == "__main__":
    if "--build" in sys.argv:
        construir_v3()
    else:
        generar_propuesta()
