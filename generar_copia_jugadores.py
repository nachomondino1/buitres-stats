"""
Genera una copia "para compartir con el plantel" a partir del Excel maestro
de Los Buitres, sacando lo que no debería ver el equipo.

Uso:
    python generar_copia_jugadores.py <archivo_maestro.xlsx> <archivo_salida.xlsx>

Qué hace (100% automático):
    - Vacía la columna "Analisis" de la hoja "Logs Partidos" (valor, hyperlink
      y comentarios) y la oculta. Ahí es donde vive el link a tu doc de
      análisis y cualquier comentario suelto tipo notas tácticas.

Qué NO puede hacer solo (te lo avisa, no lo edita):
    - Revisar la columna "Observaciones" (O) en busca de comentarios tipo
      "quién pagó y quién no" u otro señalamiento nominal. Es texto libre,
      no hay forma confiable de detectar automáticamente qué es "critica" y
      qué no. El script te imprime las filas sospechosas para que decidas vos.

Nota sobre el recálculo: este archivo usa funciones dinámicas (tipo UNIQUE/
FILTER) en algunas hojas (Export data, Torneo Apertura 2026) que Excel
calcula perfecto pero que el motor de recálculo de este entorno (LibreOffice)
no soporta y puede pisar con errores falsos. Por eso este script NO fuerza
un recálculo — al abrir el archivo en Excel de verdad, se recalcula solo.
"""
import re
import sys

import openpyxl

ANALISIS_COL = 12  # columna L
OBSERVACIONES_COL = 15  # columna O
FLAG_PATTERNS = re.compile(r"pag[oó]|✅|❌|🟡|debe|deuda", re.IGNORECASE)


def limpiar_columna_analisis(ws, max_row=200):
    for r in range(1, max_row + 1):
        cell = ws.cell(row=r, column=ANALISIS_COL)
        if cell.value is not None or cell.hyperlink is not None or cell.comment is not None:
            cell.value = None
            cell.hyperlink = None
            cell.comment = None
    ws.column_dimensions[ws.cell(row=1, column=ANALISIS_COL).column_letter].hidden = True


def revisar_observaciones(ws, max_row=200):
    sospechosas = []
    for r in range(1, max_row + 1):
        cell = ws.cell(row=r, column=OBSERVACIONES_COL)
        if isinstance(cell.value, str) and FLAG_PATTERNS.search(cell.value):
            sospechosas.append((r, cell.value))
    return sospechosas


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)

    origen, destino = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(origen, data_only=False)

    if "Logs Partidos" not in wb.sheetnames:
        print('No encontré la hoja "Logs Partidos" en el archivo. ¿Es el maestro correcto?')
        sys.exit(1)

    ws = wb["Logs Partidos"]
    limpiar_columna_analisis(ws)
    wb.save(destino)

    print(f"Listo. Copia segura guardada en: {destino}")
    print("(No se forzó recálculo — se recalcula solo al abrir en Excel real)\n")

    sospechosas = revisar_observaciones(ws)
    if sospechosas:
        print("⚠ Revisá a mano estas filas de 'Observaciones' antes de mandarlo")
        print("  (mencionan pagos o tienen ✅/❌/🟡 — pueden nombrar a alguien):")
        for fila, texto in sospechosas:
            print(f"  - O{fila}: {texto[:80]}{'...' if len(texto) > 80 else ''}")
    else:
        print("No encontré filas sospechosas en 'Observaciones'. Igual date una vuelta rápida antes de mandarlo.")


if __name__ == "__main__":
    main()
